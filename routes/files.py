import json
import mimetypes
from contextlib import contextmanager
from pathlib import Path

from fastapi import APIRouter, Depends, File, Form, HTTPException, Query, UploadFile
from fastapi.responses import FileResponse
from pydantic import BaseModel
from sqlalchemy import and_, or_
from sqlalchemy.orm import Session as DbSession
from starlette.concurrency import run_in_threadpool

from core.database import FileComment, FileTag, get_db
from services import (
    file_operations,
    files_to_photos,
    fileversions,
    storage_backends,
    storage_locations,
    textindex,
    trash,
)
from services import files_store as fs

router = APIRouter(prefix="/api/files")
MAX_UPLOAD_BYTES = 100 * 1024 * 1024
UPLOAD_READ_CHUNK = 1024 * 1024


async def _read_upload_limited(file: UploadFile, *, limit: int = MAX_UPLOAD_BYTES) -> bytes:
    """Read at most the accepted payload plus one byte, then stop the request."""
    chunks = []
    total = 0
    while True:
        chunk = await file.read(min(UPLOAD_READ_CHUNK, limit - total + 1))
        if not chunk:
            return b"".join(chunks)
        total += len(chunk)
        if total > limit:
            raise HTTPException(400, "file too large (100MB max)")
        chunks.append(chunk)


_ACTIVE_MEDIA_TYPES = frozenset(
    {
        "application/javascript",
        "application/xhtml+xml",
        "application/xml",
        "image/svg+xml",
        "text/html",
        "text/javascript",
        "text/xml",
    }
)
_REMOTE_PREVIEW_MAX_BYTES = fileversions.CAP_BYTES
_REMOTE_RAW_MAX_BYTES = 512 * 1024 * 1024


def _is_active_media_type(media_type: str) -> bool:
    value = media_type.lower()
    return value in _ACTIVE_MEDIA_TYPES or value.endswith("+xml")


class TemporaryFileResponse(FileResponse):
    """Serve a materialized remote file and remove it even if delivery aborts."""

    def __init__(self, *args, cleanup_path: Path, **kwargs):
        self.cleanup_path = Path(cleanup_path)
        super().__init__(*args, **kwargs)

    async def __call__(self, scope, receive, send):
        try:
            await super().__call__(scope, receive, send)
        finally:
            storage_backends.remove_temporary(self.cleanup_path)


def _norm_tags(tags) -> list[str]:
    """lowercase, trim, drop blanks, dedup (order-stable)."""
    seen, out = set(), []
    for t in tags or []:
        t = str(t).strip().lower()
        if t and t not in seen:
            seen.add(t)
            out.append(t)
    return out


def _location(db, location_id: str | None = None):
    try:
        return storage_locations.require_default_browsable(db, location_id)
    except LookupError as exc:
        raise HTTPException(404, str(exc)) from exc
    except RuntimeError as exc:
        raise HTTPException(409, str(exc)) from exc


def _root(location):
    try:
        root = storage_locations.local_root(location)
    except RuntimeError as exc:
        raise HTTPException(409, str(exc)) from exc
    if not root.exists() or not root.is_dir():
        raise HTTPException(409, "storage location is unavailable")
    return root


def _writable(location):
    if location.access != "managed":
        raise HTTPException(409, "storage location is read-only")
    return _root(location)


def _run_file_operation(
    db,
    *,
    action: str,
    location_id: str,
    source_path: str,
    destination_path: str = "",
):
    try:
        row = file_operations.enqueue(
            db,
            action=action,
            source_location_id=location_id,
            source_path=source_path,
            destination_location_id=(location_id if destination_path else None),
            destination_path=destination_path,
        )
        return file_operations.run(db, row)
    except LookupError as exc:
        raise HTTPException(404, str(exc)) from exc
    except (RuntimeError, ValueError, file_operations.FileOperationError) as exc:
        raise HTTPException(409, str(exc)) from exc


def _identity_path(path: str) -> str:
    try:
        return storage_locations.normalize_path(path)
    except ValueError as exc:
        raise HTTPException(400, str(exc)) from exc


@contextmanager
def _claimed_direct_mutation(db, location, path: str):
    try:
        with file_operations.direct_mutation_claim(db, location, path):
            yield
    except file_operations.FileOperationError as exc:
        raise HTTPException(409, str(exc)) from exc


def _remote_file_size(metadata: dict) -> int | None:
    value = metadata.get("size")
    if value is None:
        return None
    if not isinstance(value, int) or isinstance(value, bool) or value < 0:
        raise storage_backends.StorageBackendError("remote file size is invalid")
    return value


@contextmanager
def _materialized_preview(location, normalized: str):
    remote = location.kind != "local"
    path = None
    try:
        if remote:
            path = storage_backends.temporary_path(location, normalized)
            storage_backends.download(
                location,
                normalized,
                path,
                max_bytes=_REMOTE_PREVIEW_MAX_BYTES,
            )
        else:
            path = storage_backends.local_path(location, normalized)
        yield path
    except storage_backends.StorageNotFoundError as exc:
        raise HTTPException(404, str(exc)) from exc
    except storage_backends.StorageBackendError as exc:
        raise HTTPException(409, str(exc)) from exc
    finally:
        if remote and path is not None:
            storage_backends.remove_temporary(path)


def _legacy_or_identity(Model, path: str, location_id: str):
    normalized = _identity_path(path)
    legacy_values = {str(path or "").strip("/"), normalized}
    return and_(
        Model.location_id == location_id,
        or_(
            Model.normalized_path == normalized,
            and_(
                or_(Model.normalized_path.is_(None), Model.normalized_path == ""),
                Model.path.in_(legacy_values),
            ),
        ),
    )


def _tag_map(db, paths=None, location_id="default-local") -> dict:
    """path → {tags, color, starred}, to decorate listings cheaply. bounded to `paths` when given
    so listing a folder doesn't scan the entire FileTag table."""
    q = db.query(FileTag).filter(FileTag.location_id == location_id)
    if paths is not None:
        normalized = [_identity_path(path) for path in paths]
        q = q.filter(
            or_(
                FileTag.normalized_path.in_(normalized),
                and_(
                    or_(FileTag.normalized_path.is_(None), FileTag.normalized_path == ""),
                    FileTag.path.in_(paths),
                ),
            )
        )
    out = {}
    for r in q.all():
        out[r.normalized_path or _identity_path(r.path)] = {
            "tags": [t for t in (r.tags or "").split(",") if t],
            "color": r.color or "",
            "starred": bool(r.starred),
        }
    return out


def _decorate(items, tmap, location_id="default-local"):
    for it in items:
        normalized = _identity_path(it["path"])
        meta = tmap.get(normalized)
        it["location_id"] = location_id
        it["normalized_path"] = normalized
        it["tags"] = meta["tags"] if meta else []
        it["color"] = meta["color"] if meta else ""
        it["starred"] = meta["starred"] if meta else False
    return items


def _comment_counts(db, paths=None, location_id="default-local") -> dict:
    """path → number of comments (roots + replies), to badge listings (6c). bounded to `paths`."""
    q = db.query(FileComment.path, FileComment.normalized_path).filter(
        FileComment.location_id == location_id
    )
    if paths is not None:
        normalized = [_identity_path(path) for path in paths]
        q = q.filter(
            or_(
                FileComment.normalized_path.in_(normalized),
                and_(
                    or_(
                        FileComment.normalized_path.is_(None),
                        FileComment.normalized_path == "",
                    ),
                    FileComment.path.in_(paths),
                ),
            )
        )
    out = {}
    for p, normalized in q.all():
        p = normalized or _identity_path(p)
        out[p] = out.get(p, 0) + 1
    return out


@router.get("/list")
def list_files(
    path: str = Query(""),
    sort: str = Query("name"),
    order: str = Query(""),
    location_id: str | None = Query(None),
    db: DbSession = Depends(get_db),
):
    location = _location(db, location_id)
    normalized_path = _identity_path(path)
    try:
        d = storage_backends.listdir(location, normalized_path, sort=sort, order=order)
    except storage_backends.StorageNotFoundError as e:
        raise HTTPException(404, str(e)) from e
    except (ValueError, storage_backends.StorageBackendError) as e:
        raise HTTPException(409, str(e)) from e
    paths = [it["path"] for it in d["items"]]
    _decorate(d["items"], _tag_map(db, paths, location.id), location.id)
    cc = _comment_counts(db, paths, location.id)
    for it in d["items"]:
        it["comments"] = cc.get(it["normalized_path"], 0)
    d["location_id"] = location.id
    d["path"] = normalized_path
    return d


class TagBody(BaseModel):
    tags: list[str] | None = None
    color: str | None = None


@router.get("/tags")
def get_tags(
    path: str = Query(...),
    location_id: str | None = Query(None),
    db: DbSession = Depends(get_db),
):
    location = _location(db, location_id)
    normalized = _identity_path(path)
    r = db.query(FileTag).filter(_legacy_or_identity(FileTag, path, location.id)).first()
    if not r:
        return {
            "path": normalized,
            "location_id": location.id,
            "tags": [],
            "color": "",
        }
    return {
        "path": normalized,
        "location_id": location.id,
        "tags": [t for t in (r.tags or "").split(",") if t],
        "color": r.color or "",
    }


@router.put("/tags")
def set_tags(
    path: str = Query(...),
    body: TagBody = None,
    location_id: str | None = Query(None),
    db: DbSession = Depends(get_db),
):
    body = body or TagBody()
    location = _location(db, location_id)
    normalized = _identity_path(path)
    if not normalized:
        raise HTTPException(400, "path required")
    with _claimed_direct_mutation(db, location, normalized):
        r = db.query(FileTag).filter(_legacy_or_identity(FileTag, path, location.id)).first()
        new_tags = (
            ",".join(_norm_tags(body.tags)) if body.tags is not None else (r.tags if r else "")
        )
        new_color = body.color.strip() if body.color is not None else (r.color if r else "")
        starred = bool(r.starred) if r else False
        # Empty metadata is deleted rather than stored as scan noise.
        if not (new_tags or "").strip() and not (new_color or "").strip() and not starred:
            if r:
                db.delete(r)
                db.commit()
            result = {
                "path": normalized,
                "location_id": location.id,
                "tags": [],
                "color": "",
            }
        else:
            if not r:
                r = FileTag(
                    path=normalized,
                    location_id=location.id,
                    normalized_path=normalized,
                )
                db.add(r)
            else:
                r.path = normalized
                r.location_id = location.id
                r.normalized_path = normalized
            r.tags = new_tags
            r.color = new_color
            db.commit()
            result = {
                "path": normalized,
                "location_id": location.id,
                "tags": [t for t in (r.tags or "").split(",") if t],
                "color": r.color or "",
            }
    return result


class StarBody(BaseModel):
    starred: bool = True


@router.put("/star")
def set_star(
    path: str = Query(...),
    body: StarBody = None,
    location_id: str | None = Query(None),
    db: DbSession = Depends(get_db),
):
    location = _location(db, location_id)
    normalized = _identity_path(path)
    if not normalized:
        raise HTTPException(400, "path required")
    body = body or StarBody()
    with _claimed_direct_mutation(db, location, normalized):
        r = db.query(FileTag).filter(_legacy_or_identity(FileTag, path, location.id)).first()
        if not r:
            r = FileTag(
                path=normalized,
                location_id=location.id,
                normalized_path=normalized,
            )
            db.add(r)
        else:
            r.path = normalized
            r.location_id = location.id
            r.normalized_path = normalized
        r.starred = bool(body.starred)
        db.commit()
        result = {
            "path": normalized,
            "location_id": location.id,
            "starred": bool(r.starred),
        }
    return result


@router.get("/starred")
def list_starred(location_id: str | None = Query(None), db: DbSession = Depends(get_db)):
    location = _location(db, location_id)
    rows = (
        db.query(FileTag)
        .filter(FileTag.location_id == location.id, FileTag.starred == True)  # noqa: E712
        .all()
    )
    items = []
    for row in rows:
        path = row.normalized_path or _identity_path(row.path)
        try:
            item_type = storage_backends.item(location, path).get("type", "file")
        except storage_backends.StorageBackendError:
            item_type = "file"
        items.append(
            {
                "path": path,
                "location_id": location.id,
                "tags": [tag for tag in (row.tags or "").split(",") if tag],
                "color": row.color or "",
                "type": item_type,
            }
        )
    items.sort(key=lambda x: x["path"].lower())
    return {"location_id": location.id, "items": items}


@router.get("/quota")
def quota(location_id: str | None = Query(None), db: DbSession = Depends(get_db)):
    """bytes used by the files vault + the underlying disk's total/free (6a)."""
    import shutil

    location = _location(db, location_id)
    base = _root(location)
    used = 0
    for p in base.rglob("*"):
        try:
            if p.is_file():
                used += p.stat().st_size
        except OSError:
            pass
    try:
        du = shutil.disk_usage(str(base))
        total, free = du.total, du.free
    except OSError:
        total = free = 0
    return {"location_id": location.id, "used": used, "total": total, "free": free}


@router.get("/duplicates")
def duplicates(location_id: str | None = Query(None), db: DbSession = Depends(get_db)):
    """group files with identical content (exact SHA-256) — exact dedup only (6b)."""
    import hashlib
    from collections import defaultdict

    location = _location(db, location_id)
    base = _root(location)
    # bucket by size first: identical content => identical size, so a file with a unique size
    # can't have a dup and never needs to be read/hashed (was hashing every byte of every file).
    by_size = defaultdict(list)
    for p in base.rglob("*"):
        try:
            if p.is_file() and (sz := p.stat().st_size):
                by_size[sz].append(p)
        except OSError:
            pass
    groups = defaultdict(list)
    for sz, plist in by_size.items():
        if len(plist) < 2:
            continue
        for p in plist:
            try:
                hsh = None
                with p.open("rb") as f:
                    for chunk in iter(lambda: f.read(1024 * 1024), b""):
                        if hsh is None:
                            hsh = hashlib.sha256(chunk)
                        else:
                            hsh.update(chunk)
                h = hsh.hexdigest() if hsh else hashlib.sha256().hexdigest()
            except OSError:
                continue
            groups[h].append({"path": str(p.relative_to(base)).replace("\\", "/"), "size": sz})
    out = [
        {"hash": h, "size": items[0]["size"], "paths": [i["path"] for i in items]}
        for h, items in groups.items()
        if len(items) > 1
    ]
    out.sort(key=lambda g: -len(g["paths"]))
    return {"location_id": location.id, "groups": out}


@router.get("/preview")
def preview(
    path: str = Query(...),
    location_id: str | None = Query(None),
    db: DbSession = Depends(get_db),
):
    """office / text preview (6b). docx → text (python-docx); xlsx → openpyxl if present; txt/md → raw."""
    location = _location(db, location_id)
    normalized = _identity_path(path)
    with _materialized_preview(location, normalized) as p:
        if not p.is_file():
            raise HTTPException(404, "not found")
        return _preview_payload(p)


def _preview_payload(p: Path) -> dict:
    ext = p.suffix.lower()
    if ext == ".docx":
        try:
            import html as _html

            from docx import Document

            doc = Document(str(p))
            paras = [para.text for para in doc.paragraphs]
            text = "\n".join(paras)
            html = "".join(f"<p>{_html.escape(t)}</p>" if t.strip() else "<br>" for t in paras)
            return {"kind": "docx", "text": text, "html": html}
        except Exception as e:
            return {"kind": "docx", "error": f"couldn't read docx: {str(e)[:120]}"}
    if ext == ".xlsx":
        try:
            import openpyxl

            wb = openpyxl.load_workbook(str(p), read_only=True, data_only=True)
            try:
                rows = [
                    [("" if c is None else str(c)) for c in row]
                    for row in wb.active.iter_rows(values_only=True)
                ]
            finally:
                wb.close()  # read_only mode holds the file handle open until close()
            return {"kind": "xlsx", "rows": rows[:200]}
        except ImportError:
            return {"kind": "xlsx", "error": "openpyxl not installed (pip install openpyxl)"}
        except Exception as e:
            return {"kind": "xlsx", "error": f"couldn't read xlsx: {str(e)[:120]}"}
    if ext in (".txt", ".md", ".markdown", ".csv", ".log", ".json"):
        try:
            return {"kind": "text", "text": p.read_text("utf-8", errors="replace")[:20000]}
        except OSError:
            return {"kind": "text", "error": "couldn't read file"}
    return {"kind": "unsupported"}


@router.get("/activity")
def activity(
    days: int = 30,
    limit: int = 100,
    location_id: str | None = Query(None),
    db: DbSession = Depends(get_db),
):
    """recent file changes, newest first (6b)."""
    import time

    location = _location(db, location_id)
    base = _root(location)
    cutoff = time.time() - days * 86400
    out = []
    for p in base.rglob("*"):
        try:
            if not p.is_file():
                continue
            st = p.stat()
            if st.st_mtime < cutoff:
                continue
            out.append(
                {
                    "path": str(p.relative_to(base)).replace("\\", "/"),
                    "mtime": st.st_mtime,
                    "size": st.st_size,
                }
            )
        except OSError:
            pass
    out.sort(key=lambda x: -x["mtime"])
    return {"location_id": location.id, "items": out[:limit]}


# ── file comments (6c) — threaded like DocComment, keyed on a files-relative path ──
class CommentBody(BaseModel):
    path: str = ""
    location_id: str | None = None
    body: str = ""
    author: str = "me"
    parent_id: str | None = None


def _cdict(c):
    return {
        "id": c.id,
        "path": c.normalized_path or _identity_path(c.path),
        "location_id": c.location_id,
        "body": c.body or "",
        "author": c.author or "me",
        "parent_id": c.parent_id,
        "resolved": bool(c.resolved),
        "created_at": c.created_at.isoformat() if c.created_at else None,
    }


@router.get("/comments")
def list_comments(
    path: str = Query(...),
    location_id: str | None = Query(None),
    db: DbSession = Depends(get_db),
):
    location = _location(db, location_id)
    rows = (
        db.query(FileComment)
        .filter(_legacy_or_identity(FileComment, path, location.id))
        .order_by(FileComment.created_at.asc())
        .all()
    )
    threads = []
    for root in [c for c in rows if c.parent_id is None]:
        t = _cdict(root)
        t["replies"] = [_cdict(c) for c in rows if c.parent_id == root.id]
        threads.append(t)
    return {"threads": threads}


@router.post("/comments")
def add_comment(body: CommentBody, db: DbSession = Depends(get_db)):
    """create a thread root (needs path) or a reply (needs parent_id)."""
    if not (body.body or "").strip():
        raise HTTPException(400, "comment body required")
    if body.parent_id:
        parent = db.query(FileComment).filter_by(id=body.parent_id).first()
        if not parent:
            raise HTTPException(404, "parent comment not found")
        root = (
            parent
            if parent.parent_id is None
            else db.query(FileComment).filter_by(id=parent.parent_id).first()
        )
        if not root:
            raise HTTPException(404, "parent comment not found")
        location = _location(db, root.location_id)
        if body.location_id is not None:
            requested_location = _location(db, body.location_id)
            if requested_location.id != location.id:
                raise HTTPException(404, "parent comment not found")
        normalized = root.normalized_path or _identity_path(root.path)
        root_id = root.id
        with _claimed_direct_mutation(db, location, normalized):
            root = db.get(FileComment, root_id)
            current_path = root.normalized_path or _identity_path(root.path) if root else ""
            if not root or root.location_id != location.id or current_path != normalized:
                raise HTTPException(409, "file comment changed; retry")
            c = FileComment(
                path=normalized,
                location_id=location.id,
                normalized_path=normalized,
                body=body.body.strip(),
                author=body.author or "me",
                parent_id=root.id,
            )
            db.add(c)
            db.commit()
            db.refresh(c)
    else:
        location = _location(db, body.location_id)
        if not (body.path or "").strip():
            raise HTTPException(400, "path required")
        normalized = _identity_path(body.path)
        with _claimed_direct_mutation(db, location, normalized):
            c = FileComment(
                path=normalized,
                location_id=location.id,
                normalized_path=normalized,
                body=body.body.strip(),
                author=body.author or "me",
            )
            db.add(c)
            db.commit()
            db.refresh(c)
    return _cdict(c)


@router.post("/comments/{cid}/resolve")
def resolve_comment(cid: str, db: DbSession = Depends(get_db)):
    c = db.query(FileComment).filter_by(id=cid).first()
    if not c:
        raise HTTPException(404, "comment not found")
    root = c if c.parent_id is None else db.query(FileComment).filter_by(id=c.parent_id).first()
    if not root:
        raise HTTPException(404, "comment not found")
    location = _location(db, root.location_id)
    normalized = root.normalized_path or _identity_path(root.path)
    root_id = root.id
    with _claimed_direct_mutation(db, location, normalized):
        root = db.get(FileComment, root_id)
        current_path = root.normalized_path or _identity_path(root.path) if root else ""
        if not root or root.location_id != location.id or current_path != normalized:
            raise HTTPException(409, "file comment changed; retry")
        root.resolved = not bool(root.resolved)
        db.commit()
        result = {"id": root.id, "resolved": bool(root.resolved)}
    return result


@router.delete("/comments/{cid}")
def delete_comment(cid: str, db: DbSession = Depends(get_db)):
    """delete a comment; deleting a root also drops its replies."""
    c = db.query(FileComment).filter_by(id=cid).first()
    if not c:
        raise HTTPException(404, "comment not found")
    location = _location(db, c.location_id)
    normalized = c.normalized_path or _identity_path(c.path)
    with _claimed_direct_mutation(db, location, normalized):
        c = db.get(FileComment, cid)
        current_path = c.normalized_path or _identity_path(c.path) if c else ""
        if not c or c.location_id != location.id or current_path != normalized:
            raise HTTPException(409, "file comment changed; retry")
        if c.parent_id is None:
            db.query(FileComment).filter_by(parent_id=c.id).delete()
        db.delete(c)
        db.commit()
    return {"ok": True}


@router.get("/by-tag")
def files_by_tag(
    tag: str = Query(...),
    location_id: str | None = Query(None),
    db: DbSession = Depends(get_db),
):
    location = _location(db, location_id)
    t = tag.strip().lower()
    out = []
    for r in db.query(FileTag).filter(FileTag.location_id == location.id).all():
        tags = [x for x in (r.tags or "").split(",") if x]
        if t in tags:
            out.append(
                {
                    "path": r.normalized_path or _identity_path(r.path),
                    "location_id": location.id,
                    "tags": tags,
                    "color": r.color or "",
                }
            )
    out.sort(key=lambda x: x["path"].lower())
    return {"tag": t, "location_id": location.id, "items": out}


@router.get("/tags/all")
def all_tags(location_id: str | None = Query(None), db: DbSession = Depends(get_db)):
    location = _location(db, location_id)
    tags, colors = set(), {}
    for r in db.query(FileTag).filter(FileTag.location_id == location.id).all():
        for x in (r.tags or "").split(","):
            if x:
                tags.add(x)
        if r.color:
            colors[r.normalized_path or _identity_path(r.path)] = r.color
    return {"location_id": location.id, "tags": sorted(tags), "colors": colors}


@router.get("/search")
def search_files(
    q: str = Query(...),
    limit: int = 100,
    location_id: str | None = Query(None),
    db: DbSession = Depends(get_db),
):
    location = _location(db, location_id)
    if location.kind == "local":
        result = fs.search(q, limit, root=_root(location))
    else:
        result = {
            "query": q,
            "results": textindex.search_file_location(db, q, location.id, limit),
        }
    result["location_id"] = location.id
    for item in result["results"]:
        item["location_id"] = location.id
        item["normalized_path"] = _identity_path(item.get("path", ""))
    return result


@router.get("/smart/{kind}")
def smart_folder(
    kind: str,
    days: int = 30,
    limit: int = 200,
    location_id: str | None = Query(None),
    db: DbSession = Depends(get_db),
):
    location = _location(db, location_id)
    try:
        result = fs.smart(kind, days=days, limit=limit, root=_root(location))
    except ValueError as e:
        raise HTTPException(400, str(e))
    result["location_id"] = location.id
    for item in result.get("items", []):
        item["location_id"] = location.id
        item["normalized_path"] = _identity_path(item.get("path", ""))
    return result


@router.get("/read")
def read_file(
    path: str = Query(...),
    location_id: str | None = Query(None),
    db: DbSession = Depends(get_db),
):
    location = _location(db, location_id)
    normalized = _identity_path(path)
    try:
        result = storage_backends.read_text(location, normalized)
    except storage_backends.StorageNotFoundError as e:
        raise HTTPException(404, str(e)) from e
    except (ValueError, storage_backends.StorageBackendError) as e:
        raise HTTPException(409, str(e)) from e
    if isinstance(result, dict):
        result["location_id"] = location.id
        result["normalized_path"] = normalized
    return result


@router.get("/raw")
def raw_file(
    path: str = Query(...),
    download: bool = False,
    location_id: str | None = Query(None),
    db: DbSession = Depends(get_db),
):
    location = _location(db, location_id)
    normalized = _identity_path(path)
    remote = location.kind != "local"
    if not remote:
        try:
            p = storage_backends.local_path(location, normalized)
        except storage_backends.StorageNotFoundError as e:
            raise HTTPException(404, str(e)) from e
        except storage_backends.StorageBackendError as e:
            raise HTTPException(409, str(e)) from e
    else:
        p = storage_backends.temporary_path(location, normalized)
        try:
            storage_backends.download(
                location,
                normalized,
                p,
                max_bytes=_REMOTE_RAW_MAX_BYTES,
            )
        except storage_backends.StorageNotFoundError as e:
            storage_backends.remove_temporary(p)
            raise HTTPException(404, str(e)) from e
        except storage_backends.StorageBackendError as e:
            storage_backends.remove_temporary(p)
            raise HTTPException(409, str(e)) from e
    if not p.is_file():
        if remote:
            storage_backends.remove_temporary(p)
        raise HTTPException(404)
    # Inline safe media by default so images preview. Active document types can
    # execute under the Alles origin, so they must always download as inert data.
    response_type = TemporaryFileResponse if remote else FileResponse
    media_type = mimetypes.guess_type(normalized)[0] or "application/octet-stream"
    active_content = _is_active_media_type(media_type)
    headers = {"X-Content-Type-Options": "nosniff"}
    options = {"media_type": media_type, "headers": headers}
    if download or active_content:
        options.update(
            {
                "filename": Path(normalized).name or "file",
                "content_disposition_type": "attachment",
            }
        )
    if active_content:
        options["media_type"] = "application/octet-stream"
        headers["Content-Security-Policy"] = "sandbox; default-src 'none'"
    if remote:
        options["cleanup_path"] = p
    return response_type(str(p), **options)


class MkdirBody(BaseModel):
    path: str
    location_id: str | None = None


@router.post("/mkdir")
def mkdir(body: MkdirBody, db: DbSession = Depends(get_db)):
    location = _location(db, body.location_id)
    normalized = _identity_path(body.path)
    try:
        if location.access != "managed":
            raise HTTPException(409, "storage location is read-only")
        if location.kind == "local":
            _writable(location)
        with file_operations.direct_mutation_claim(db, location, normalized):
            result = storage_backends.mkdir(location, normalized)
    except (
        ValueError,
        file_operations.FileOperationError,
        storage_backends.StorageBackendError,
    ) as e:
        raise HTTPException(409, str(e))
    result["location_id"] = location.id
    result["normalized_path"] = normalized
    return result


class FilesToPhotosBody(BaseModel):
    path: str
    location_id: str | None = None
    album_id: str = ""


@router.post("/to-photos")
def to_photos(body: FilesToPhotosBody, db: DbSession = Depends(get_db)):
    location = _location(db, body.location_id)
    try:
        return files_to_photos.import_from_files(
            db,
            location_id=location.id,
            path=body.path,
            album_id=body.album_id,
        )
    except LookupError as exc:
        raise HTTPException(404, str(exc)) from exc
    except (
        RuntimeError,
        ValueError,
        files_to_photos.FilesToPhotosError,
        storage_backends.StorageBackendError,
    ) as exc:
        raise HTTPException(409, str(exc)) from exc


class RenameBody(BaseModel):
    path: str
    to: str
    location_id: str | None = None


def _meta_models():
    from core.database import FileVersion

    return (FileTag, FileComment, FileVersion)


def _path_filter(Model, path, location_id):
    normalized = _identity_path(path)
    pre = normalized.rstrip("/") + "/"
    return and_(
        Model.location_id == location_id,
        or_(
            Model.normalized_path == normalized,
            Model.normalized_path.startswith(pre),
            and_(
                or_(Model.normalized_path.is_(None), Model.normalized_path == ""),
                or_(Model.path == path, Model.path.startswith(pre)),
            ),
        ),
    )


@router.post("/rename")
def rename(body: RenameBody, db: DbSession = Depends(get_db)):
    location = _location(db, body.location_id)
    source = _identity_path(body.path)
    destination = _identity_path(body.to)
    if location.kind == "local":
        root = _writable(location)
        try:
            if fs.abspath(destination, root=root).exists():
                raise HTTPException(400, "destination already exists")
        except ValueError as e:
            raise HTTPException(400, str(e))
    _run_file_operation(
        db,
        action="rename",
        location_id=location.id,
        source_path=source,
        destination_path=destination,
    )
    return {
        "ok": True,
        "path": destination,
        "location_id": location.id,
        "normalized_path": destination,
    }


@router.delete("/delete")
def delete(
    path: str = Query(...),
    location_id: str | None = Query(None),
    db: DbSession = Depends(get_db),
):
    location = _location(db, location_id)
    normalized = _identity_path(path)
    if not normalized:
        raise HTTPException(400, "won't delete the root")
    if location.kind == "local":
        root = _writable(location)
        try:
            p = fs.abspath(normalized, root=root)
        except ValueError as e:
            raise HTTPException(400, str(e))
        if not p.exists():
            raise HTTPException(404, "not found")
    _run_file_operation(
        db,
        action="delete",
        location_id=location.id,
        source_path=normalized,
    )
    return {
        "ok": True,
        "trashed": True,
        "location_id": location.id,
        "normalized_path": normalized,
    }


@router.get("/trash")
def list_trash(location_id: str | None = Query(None), db: DbSession = Depends(get_db)):
    location = _location(db, location_id)
    items = [
        item
        for item in trash.list_items(db, kind="file")
        if (item.location_id or "default-local") == location.id
    ]
    return [
        {
            "id": it.id,
            "ref": it.normalized_path or _identity_path(it.ref),
            "location_id": location.id,
            "name": it.name,
            "type": "dir" if json.loads(it.payload or "{}").get("is_dir") else "file",
            "trashed_at": it.trashed_at.isoformat() if it.trashed_at else None,
        }
        for it in items
    ]


class TrashAction(BaseModel):
    id: str
    location_id: str | None = None


@router.post("/trash/restore")
def restore_trash(body: TrashAction, db: DbSession = Depends(get_db)):
    location = _location(db, body.location_id)
    it = trash.get(db, body.id)
    if not it or it.kind != "file" or (it.location_id or "default-local") != location.id:
        raise HTTPException(404)
    normalized = it.normalized_path or _identity_path(it.ref)
    if location.kind == "local":
        _writable(location)
    _run_file_operation(
        db,
        action="restore",
        location_id=location.id,
        source_path=body.id,
    )
    return {
        "ok": True,
        "restored": normalized,
        "location_id": location.id,
    }


@router.post("/trash/purge")
def purge_trash(location_id: str | None = None, db: DbSession = Depends(get_db)):
    location = _location(db, location_id)
    if location.access != "managed":
        raise HTTPException(409, "storage location is read-only")
    return {
        "purged": trash.purge_expired(
            db,
            kind="file",
            location_id=location.id,
        )
    }


@router.post("/upload")
async def upload(
    path: str = Form(""),
    location_id: str | None = Form(None),
    expected_etag: str = Form(""),
    file: UploadFile = File(...),
    db: DbSession = Depends(get_db),
):
    location = _location(db, location_id)
    if location.access != "managed":
        raise HTTPException(409, "storage location is read-only")
    normalized_dir = _identity_path(path)
    data = await _read_upload_limited(file)
    name = Path(file.filename).name
    if not name or set(name) <= {"."}:
        raise HTTPException(400, "invalid filename")
    rel = _identity_path(f"{normalized_dir}/{name}" if normalized_dir else name)
    try:
        with file_operations.direct_mutation_claim(db, location, rel):
            if location.kind == "local":
                root = _writable(location)
                try:
                    existing = fs.abspath(rel, root=root)
                    if existing.is_file():  # overwrite → snapshot the old content first (1e)
                        fileversions.snapshot(db, rel, existing, location.id)
                    result = fs.save_upload(normalized_dir, file.filename, data, root=root)
                except ValueError as e:
                    raise HTTPException(400, str(e))
                except OSError as e:
                    raise HTTPException(409, "upload could not be saved") from e
            else:

                def store_remote():
                    source = storage_backends.temporary_path(location, name)
                    previous = storage_backends.temporary_path(location, rel)
                    try:
                        try:
                            existing = storage_backends.item(location, rel)
                        except storage_backends.StorageNotFoundError:
                            existing = None
                        if existing is not None:
                            if existing.get("type") != "file":
                                raise HTTPException(409, "upload destination is not a file")
                            if not expected_etag:
                                raise HTTPException(
                                    409,
                                    "remote file already exists; reload before overwriting",
                                )
                            current_etag = str(existing.get("etag") or "")
                            if not current_etag or current_etag != expected_etag:
                                raise HTTPException(409, "remote file changed")
                            current_size = _remote_file_size(existing)
                            if current_size is not None and current_size > fileversions.CAP_BYTES:
                                raise HTTPException(
                                    409,
                                    "remote file is too large to version before overwrite",
                                )
                            storage_backends.download(
                                location,
                                rel,
                                previous,
                                expected_etag=expected_etag,
                                expected_size=current_size,
                                max_bytes=fileversions.CAP_BYTES,
                            )
                            fileversions.snapshot(db, rel, previous, location.id)
                        source.write_bytes(data)
                        stored = storage_backends.upload(
                            location,
                            rel,
                            source,
                            expected_etag=expected_etag,
                        )
                        stored.update({"ok": True, "name": name, "path": rel})
                        return stored
                    except storage_backends.StorageBackendError as e:
                        raise HTTPException(409, str(e)) from e
                    finally:
                        storage_backends.remove_temporary(source)
                        storage_backends.remove_temporary(previous)

                result = await run_in_threadpool(store_remote)
    except file_operations.FileOperationError as exc:
        raise HTTPException(409, str(exc)) from exc
    result["location_id"] = location.id
    result["normalized_path"] = _identity_path(result["path"])
    return result


@router.get("/versions")
def list_versions(
    path: str = Query(...),
    location_id: str | None = Query(None),
    db: DbSession = Depends(get_db),
):
    location = _location(db, location_id)
    normalized = _identity_path(path)
    return [
        {
            "id": v.id,
            "location_id": location.id,
            "normalized_path": normalized,
            "size": v.size,
            "sha": v.sha[:12],
            "created_at": v.created_at.isoformat() if v.created_at else None,
        }
        for v in fileversions.list_versions(db, normalized, location.id)
    ]


class RestoreVersion(BaseModel):
    path: str
    id: str
    location_id: str | None = None


@router.post("/versions/restore")
def restore_version(body: RestoreVersion, db: DbSession = Depends(get_db)):
    location = _location(db, body.location_id)
    normalized = _identity_path(body.path)
    try:
        with file_operations.direct_mutation_claim(db, location, normalized):
            version = fileversions.get(db, body.id)
            if (
                not version
                or version.location_id != location.id
                or (version.normalized_path or _identity_path(version.path)) != normalized
            ):
                raise HTTPException(404)
            if location.kind == "local":
                root = _writable(location)
                # snapshot the current content first so a restore is itself undoable
                cur = fs.abspath(normalized, root=root)
                if cur.is_file():
                    fileversions.snapshot(db, normalized, cur, location.id)
                if not fileversions.restore(db, body.id, cur):
                    raise HTTPException(404)
            else:
                if location.access != "managed":
                    raise HTTPException(409, "storage location is read-only")
                current = storage_backends.temporary_path(location, normalized)
                restored = storage_backends.temporary_path(location, normalized)
                try:
                    if not fileversions.restore(db, body.id, restored):
                        raise HTTPException(404)
                    try:
                        metadata = storage_backends.item(location, normalized)
                    except storage_backends.StorageNotFoundError:
                        metadata = None
                    current_etag = None
                    if metadata is not None:
                        if metadata.get("type") != "file":
                            raise HTTPException(409, "version destination is not a file")
                        current_etag = str(metadata.get("etag") or "")
                        if not current_etag:
                            raise HTTPException(409, "remote file does not expose a stable ETag")
                        current_size = _remote_file_size(metadata)
                        if current_size is not None and current_size > fileversions.CAP_BYTES:
                            raise HTTPException(409, "current file is too large to version safely")
                        storage_backends.download(
                            location,
                            normalized,
                            current,
                            expected_etag=current_etag,
                            expected_size=current_size,
                            max_bytes=fileversions.CAP_BYTES,
                        )
                        fileversions.snapshot(db, normalized, current, location.id)
                    storage_backends.upload(
                        location,
                        normalized,
                        restored,
                        expected_etag=current_etag,
                    )
                except storage_backends.StorageBackendError as exc:
                    raise HTTPException(409, str(exc)) from exc
                finally:
                    storage_backends.remove_temporary(current)
                    storage_backends.remove_temporary(restored)
    except file_operations.FileOperationError as exc:
        raise HTTPException(409, str(exc)) from exc
    return {"ok": True, "restored": normalized, "location_id": location.id}
