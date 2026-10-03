"""Local import preview and the existing durable create path preserve source content."""

import io
import uuid
from unittest import mock

from docx import Document

from services import document_safety, vault_md
from tests._client import VaultApiTest


class DocumentImportTests(VaultApiTest):
    def preview(self, name, content):
        return self.client.post("/api/vault-md/import-preview", files={"file": (name, content)})

    def test_text_preview_and_durable_create_preserve_exact_bytes(self):
        for extension in ("md", "markdown", "txt"):
            with self.subTest(extension=extension):
                source = "    indented code\r\n\r\nnext  \r\n\r\n"
                before = vault_md.tree()
                response = self.preview(f"original.{extension}", source.encode())
                self.assertEqual(response.status_code, 200, response.text)
                data = response.json()
                self.assertEqual(data["content"], source)
                self.assertEqual(data["name"], "original")
                self.assertEqual(vault_md.tree(), before)
                body = {
                    "path": f"{extension}.md",
                    "content": data["content"],
                    "unique": True,
                    "request_id": str(uuid.uuid4()),
                }
                first = self.client.post("/api/vault-md/file", json=body)
                retry = self.client.post("/api/vault-md/file", json=body)
                self.assertEqual(first.status_code, 200, first.text)
                self.assertEqual(retry.json()["path"], first.json()["path"])
                self.assertEqual(vault_md._safe(first.json()["path"]).read_bytes(), source.encode())

    def test_bad_encoding_or_empty_extraction_never_creates_a_document(self):
        for name, source in [
            ("bad.md", b"\xff"),
            ("bad.txt", b"\xfe"),
            ("bad.html", b"\xff"),
            ("empty.md", b" \r\n\t"),
            ("empty.html", b"<html><body></body></html>"),
        ]:
            with self.subTest(name=name):
                before = vault_md.tree()
                with mock.patch.object(
                    document_safety,
                    "create_document",
                    side_effect=AssertionError("preview wrote a document"),
                ):
                    response = self.preview(name, source)
                self.assertEqual(response.status_code, 400, response.text)
                self.assertEqual(vault_md.tree(), before)

    def test_pdf_resource_limit_returns_413_without_writing(self):
        from services.doc_import import DocumentTooLarge

        before = vault_md.tree()
        with mock.patch(
            "services.doc_import._pdf_to_md", side_effect=DocumentTooLarge("split the document")
        ):
            response = self.preview("local.pdf", b"local fixture")
        self.assertEqual(response.status_code, 413, response.text)
        self.assertIn("split the document", response.json()["detail"])
        self.assertEqual(vault_md.tree(), before)

    def test_docx_preserves_paragraph_table_paragraph_order(self):
        doc = Document()
        doc.add_paragraph("before table")
        table = doc.add_table(rows=2, cols=2)
        for row, values in zip(
            table.rows, [("column one", "column two"), ("cell one", "cell two")]
        ):
            for cell, value in zip(row.cells, values):
                cell.text = value
        doc.add_paragraph("after table")
        data = io.BytesIO()
        doc.save(data)
        response = self.preview("ordered.docx", data.getvalue())
        self.assertEqual(response.status_code, 200, response.text)
        text = response.json()["content"]
        self.assertLess(text.index("before table"), text.index("cell one"))
        self.assertLess(text.index("cell two"), text.index("after table"))
        self.assertTrue(response.json()["warning"])

    def test_invalid_document_unsupported_type_and_size_fail_before_write(self):
        for name, source in [
            ("broken.docx", b"not a document"),
            ("archive.zip", b"no"),
            ("large.txt", b"x" * (8 * 1024 * 1024 + 1)),
        ]:
            with self.subTest(name=name):
                before = vault_md.tree()
                response = self.preview(name, source)
                self.assertIn(response.status_code, (400, 413), response.text)
                self.assertEqual(vault_md.tree(), before)

    def test_extraction_size_limit_and_failed_conversion_do_not_return_partial_text(self):
        for result in (
            {"name": "large", "content": "x" * (1024 * 1024 + 1)},
            {"name": "empty", "content": ""},
        ):
            with mock.patch("services.doc_import.import_document", return_value=result):
                response = self.preview("example.docx", b"local fixture")
            self.assertIn(response.status_code, (400, 413), response.text)
        with mock.patch(
            "services.doc_import.import_document",
            side_effect=ValueError("could not extract all pages"),
        ):
            response = self.preview("example.pdf", b"local fixture")
        self.assertEqual(response.status_code, 400)
        self.assertIn("could not extract all pages", response.json()["detail"])
