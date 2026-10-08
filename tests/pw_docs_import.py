"""Owned local document import: preview, acceptance, exact storage, and recovery."""

import io
import json
import os
import re
from pathlib import Path
from urllib.parse import urlparse

from browser_gate_safety import require_server_ownership
from docx import Document
from playwright.sync_api import expect, sync_playwright


def run():
    base = f"http://127.0.0.1:{os.environ['PORT']}"
    require_server_ownership(base, os.environ["ALLES_TEST_RUN_ID"])
    out = Path(os.environ["ALLES_BROWSER_ARTIFACTS"])
    rows = []
    with sync_playwright() as pw:
        browser = pw.chromium.launch()
        api = pw.request.new_context(base_url=base)
        assert api.post("/api/setup/dismiss").ok
        for width in (1440, 820, 390):
            context = browser.new_context(
                viewport={"width": width, "height": 900},
                reduced_motion="reduce",
                service_workers="block",
            )
            context.route(
                "**/*",
                lambda route: (
                    route.continue_()
                    if urlparse(route.request.url).netloc == urlparse(base).netloc
                    else route.abort()
                ),
            )
            page = context.new_page()
            page.set_default_timeout(8000)
            errors, console, writes = [], [], []
            page.on("pageerror", lambda error: errors.append(str(error)))
            page.on(
                "console",
                lambda message: console.append(message.text) if message.type == "error" else None,
            )
            page.on(
                "request",
                lambda request: (
                    writes.append(request.post_data_json)
                    if request.url == base + "/api/vault-md/file" and request.method == "POST"
                    else None
                ),
            )
            context.tracing.start(screenshots=True, snapshots=True)

            def record(name):
                rows.append(
                    {
                        "scenario_id": "docs.import." + name,
                        "profile": str(width),
                        "status": "passed",
                    }
                )
                (out / "scenarios.json").write_text(json.dumps(rows, indent=2))

            def choose(name, content):
                with page.expect_file_chooser() as chooser:
                    page.locator("#docs-import-choose").click()
                chooser.value.set_files(
                    {"name": name, "mimeType": "application/octet-stream", "buffer": content}
                )

            try:
                page.goto(base + "/?view=wiki", wait_until="networkidle")
                trigger = page.locator("#docs-import-btn:visible, #docs-empty-import:visible")
                trigger.focus()
                page.keyboard.press("Enter")
                expect(page.locator("#docs-import-choose")).to_be_focused()
                expect(page.locator("#docs-import-name")).to_be_disabled()
                source = f"    code {width}\r\n\r\nimport searchable {width}  \r\n\r\n"
                choose(f"original-{width}.md", source.encode())
                expect(page.locator("#docs-import-save")).to_be_enabled()
                assert page.locator("#docs-import-preview").text_content() == source
                assert not writes
                page.screenshot(path=str(out / f"{width}-preview.png"), full_page=True)
                for identifier in ("docs-import-choose", "docs-import-close", "docs-import-save"):
                    assert page.locator("#" + identifier).bounding_box()["height"] >= 44
                page.locator("#docs-import-close").click()
                expect(trigger).to_be_focused()
                assert not writes
                record("preview-cancel-no-write")

                trigger.click()
                for name, content, message in [
                    ("bad.txt", b"\xff", "UTF-8"),
                    ("empty.md", b" \r\n", "no readable text"),
                    (
                        "formatting.html",
                        b"<body><h1></h1><p><b> </b></p></body>",
                        "no readable text",
                    ),
                    ("broken.docx", b"not a docx", "could not read"),
                ]:
                    choose(name, content)
                    expect(page.locator("#docs-import-status")).to_contain_text(message)
                    expect(page.locator("#docs-import-save")).to_be_disabled()
                    assert not writes
                record("invalid-empty-failed-no-write")

                choose(f"original-{width}.md", source.encode())
                expect(page.locator("#docs-import-save")).to_be_enabled()
                name = f"import-{width}.md"
                page.locator("#docs-import-name").fill(name)
                assert api.post(
                    "/api/vault-md/file", data={"path": name, "content": "existing document"}
                ).ok
                replies = []

                def lose(route):
                    if route.request.method != "POST":
                        route.continue_()
                        return
                    response = route.fetch()
                    assert response.ok, response.text()
                    replies.append(response.json())
                    route.fulfill(status=503, json={"detail": "synthetic lost import response"})

                page.route(base + "/api/vault-md/file", lose)
                page.locator("#docs-import-save").click()
                expect(page.locator("#docs-import-recovery .note-retry")).to_be_visible()
                assert writes[0]["content"] == source
                path = replies[0]["path"]
                assert path != name
                assert (
                    api.get("/api/vault-md/file", params={"path": name}).json()["content"]
                    == "existing document"
                )
                page.unroute(base + "/api/vault-md/file", lose)
                page.reload(wait_until="networkidle")
                retry = page.locator("#docs-import-recovery .note-retry")
                expect(retry).to_be_visible()
                retry.focus()
                page.keyboard.press("Enter")
                opened = page.locator("#docs-import-recovery .note-saved-open")
                expect(opened).to_be_focused()
                assert writes[-1] == writes[0]
                assert (
                    api.get("/api/vault-md/file", params={"path": path}).json()["content"] == source
                )
                expect(page.locator("#docs-import-recovery")).to_contain_text(path)
                opened.click()
                expect(page.locator("#wiki-preview")).to_contain_text(f"import searchable {width}")
                assert (
                    not page.locator("#wiki-preview")
                    .locator("[onerror], [onclick], [onload]")
                    .count()
                )
                record("collision-lost-response-reload-exact-retry")

                assert any(
                    hit["path"] == path
                    for hit in api.get(
                        "/api/vault-md/grep", params={"q": f"import searchable {width}"}
                    ).json()["results"]
                )
                page.locator("#wiki-edit-btn").click()
                page.locator("#wiki-source-btn").click()
                expect(page.locator("#wiki-source")).to_have_value(source.replace("\r\n", "\n"))
                page.locator("#wiki-source").fill(source + "edited after import\n")
                page.locator("#wiki-save-btn").click()
                expect(page.locator("#wiki-save-state")).to_contain_text("saved", ignore_case=True)
                assert (
                    "edited after import"
                    in api.get("/api/vault-md/file", params={"path": path}).json()["content"]
                )
                record("open-find-edit-save")
                held_draft = []

                def hold_draft(route):
                    if route.request.method == "PUT" and not held_draft:
                        held_draft.append((route, route.fetch()))
                    else:
                        route.continue_()

                page.route(base + "/api/vault-md/safety/draft", hold_draft)
                page.locator("#wiki-source").fill(source + "edited after import\npending draft\n")
                trigger.click()
                for _ in range(100):
                    if held_draft:
                        break
                    page.wait_for_timeout(20)
                assert len(held_draft) == 1
                page.locator("#wiki-source").focus()
                route, response = held_draft[0]
                route.fulfill(response=response)
                page.unroute(base + "/api/vault-md/safety/draft", hold_draft)
                expect(page.locator("#docs-import-panel")).to_be_visible()
                expect(page.locator("#wiki-source")).to_be_focused()
                page.locator("#docs-import-close").click()
                page.locator("#wiki-save-btn").click()
                expect(page.locator("#wiki-save-state")).to_contain_text("saved", ignore_case=True)
                record("opening-import-preserves-focus-after-draft-flush")
                trigger.click()
                expect(page.locator("#docs-import-panel")).to_be_visible()
                held_draft.clear()
                page.route(base + "/api/vault-md/safety/draft", hold_draft)
                page.locator("#wiki-source").fill(
                    source + "edited after import\ncancelled import draft\n"
                )
                trigger.click()
                for _ in range(100):
                    if held_draft:
                        break
                    page.wait_for_timeout(20)
                assert len(held_draft) == 1
                page.locator("#docs-import-close").click()
                route, response = held_draft[0]
                route.fulfill(response=response)
                page.unroute(base + "/api/vault-md/safety/draft", hold_draft)
                expect(page.locator("#wiki-save-state")).to_have_text(
                    re.compile(r"^(draft kept locally|local draft · not saved)$")
                )
                # Let the draft's awaiting import handler finish before checking cancellation.
                page.evaluate(
                    "() => new Promise(resolve => requestAnimationFrame(() => requestAnimationFrame(resolve)))"
                )
                expect(page.locator("#docs-import-panel")).to_be_hidden()
                page.locator("#wiki-save-btn").click()
                expect(page.locator("#wiki-save-state")).to_contain_text("saved", ignore_case=True)
                record("cancel-invalidates-pending-import-entry")

                trigger.click()
                choose(f"consumed-{width}.md", b"completed preview fixture")
                expect(page.locator("#docs-import-save")).to_be_enabled()
                held_draft.clear()
                page.route(base + "/api/vault-md/safety/draft", hold_draft)
                page.locator("#wiki-source").fill(
                    source + "edited after import\npending while import completes\n"
                )
                trigger.click()
                for _ in range(100):
                    if held_draft:
                        break
                    page.wait_for_timeout(20)
                assert len(held_draft) == 1
                page.locator("#docs-import-save").click()
                expect(page.locator("#docs-import-panel")).to_be_hidden()
                route, response = held_draft[0]
                route.fulfill(response=response)
                page.unroute(base + "/api/vault-md/safety/draft", hold_draft)
                page.evaluate(
                    "() => new Promise(resolve => requestAnimationFrame(() => requestAnimationFrame(resolve)))"
                )
                stayed_closed = page.locator("#docs-import-panel").is_hidden()
                trigger.click()
                expect(page.locator("#docs-import-panel")).to_be_visible()
                state = {
                    "stayed_closed": stayed_closed,
                    "text": page.locator("#docs-import-preview").text_content(),
                    "name": page.locator("#docs-import-name").input_value(),
                    "status": page.locator("#docs-import-status").text_content(),
                }
                assert state == {
                    "stayed_closed": True,
                    "text": "",
                    "name": "",
                    "status": "nothing is saved until you choose import",
                }, state
                page.locator("#docs-import-close").click()
                page.locator("#wiki-save-btn").click()
                expect(page.locator("#wiki-save-state")).to_contain_text("saved", ignore_case=True)
                record("completed-save-clears-preview-and-invalidates-entry")

                page.locator("#wiki-done-btn").click()
                trigger.click()
                choose(f"pending-{width}.md", b"    same content\n\n")
                expect(page.locator("#docs-import-save")).to_be_enabled()
                page.locator("#docs-import-name").fill(f"pending-a-{width}.md")
                page.route(base + "/api/vault-md/file", lose)
                page.locator("#docs-import-save").click()
                expect(page.locator("#docs-import-recovery .note-retry")).to_be_visible()
                first_attempt_count = len(writes)
                page.locator("#docs-import-name").fill(f"pending-b-{width}.md")
                page.locator("#docs-import-save").click()
                expect(page.locator("#docs-import-status")).to_contain_text("pending note")
                assert len(writes) == first_attempt_count
                page.unroute(base + "/api/vault-md/file", lose)
                page.locator("#docs-import-recovery .note-retry").click()
                expect(page.locator("#docs-import-recovery .note-saved-open")).to_be_visible()
                expect(page.locator("#docs-import-panel")).to_be_visible()
                expect(page.locator("#docs-import-name")).to_have_value(f"pending-b-{width}.md")
                page.locator("#docs-import-close").click()
                record("pending-save-does-not-discard-new-destination")
                trigger.click()
                choose(f"recovery-focus-{width}.md", b"local recovery focus fixture")
                expect(page.locator("#docs-import-save")).to_be_enabled()
                page.route(base + "/api/vault-md/file", lose)
                page.locator("#docs-import-save").click()
                expect(page.locator("#docs-import-recovery .note-retry")).to_be_visible()
                page.unroute(base + "/api/vault-md/file", lose)
                held_retry = []

                def hold_retry(route):
                    if route.request.method == "POST":
                        held_retry.append((route, route.fetch()))
                    else:
                        route.continue_()

                page.route(base + "/api/vault-md/file", hold_retry)
                page.locator("#docs-import-save").click()
                for _ in range(100):
                    if held_retry:
                        break
                    page.wait_for_timeout(20)
                assert len(held_retry) == 1
                page.locator("#docs-import-recovery .note-retry").focus()
                expect(page.locator("#docs-import-recovery .note-retry")).to_be_focused()
                route, response = held_retry[0]
                route.fulfill(response=response)
                page.unroute(base + "/api/vault-md/file", hold_retry)
                expect(page.locator("#docs-import-panel")).to_be_hidden()
                expect(page.locator("#docs-import-recovery .note-saved-open")).to_be_focused()
                record("direct-save-retry-transfers-recovery-focus")

                trigger.click()
                exact = "    ordinary capture differs  \r\n\r\n"
                requested = f"normalized-{width}.md"
                page.route(base + "/api/vault-md/file", lose)
                page.evaluate(
                    """async ({text, path}) => {
                        const {saveNote} = await import('/static/js/note_capture.js');
                        try { await saveNote(text, path); } catch (_) {}
                    }""",
                    {"text": exact, "path": requested},
                )
                page.unroute(base + "/api/vault-md/file", lose)
                choose(requested, exact.encode())
                expect(page.locator("#docs-import-save")).to_be_enabled()
                page.locator("#docs-import-save").click()
                expect(page.locator("#docs-import-status")).to_contain_text("pending note")
                page.locator("#docs-import-recovery .note-retry").click()
                expect(page.locator("#docs-import-recovery .note-saved-open")).to_be_visible()
                expect(page.locator("#docs-import-panel")).to_be_visible()
                assert page.locator("#docs-import-preview").text_content() == exact
                assert (
                    api.get("/api/vault-md/file", params={"path": requested}).json()["content"]
                    == exact.strip() + "\n"
                )
                page.locator("#docs-import-save").click()
                expect(page.locator("#docs-import-panel")).to_be_hidden()
                assert writes[-1]["content"] == exact
                expect(page.locator("#docs-import-recovery")).to_contain_text(
                    f"normalized-{width} 2.md"
                )
                assert (
                    api.get(
                        "/api/vault-md/file", params={"path": f"normalized-{width} 2.md"}
                    ).json()["content"]
                    == exact
                )
                record("ordinary-recovery-retains-different-exact-payload")

                trigger.click()
                choose(
                    f"html-{width}.html",
                    b"<html><head><title>metadata</title><p>Report</p><script>hidden()</script></html>",
                )
                expect(page.locator("#docs-import-save")).to_be_enabled()
                assert page.locator("#docs-import-preview").text_content() == "Report\n"
                page.locator("#docs-import-save").click()
                expect(page.locator("#docs-import-recovery .note-saved-open")).to_be_visible()
                page.locator("#docs-import-recovery .note-saved-open").click()
                expect(page.locator("#wiki-path")).to_have_text(f"html-{width}.md")
                expect(page.locator("#wiki-preview")).to_have_text("Report")
                assert (
                    api.get("/api/vault-md/file", params={"path": f"html-{width}.md"}).json()[
                        "content"
                    ]
                    == "Report\n"
                )
                record("implicit-head-html-save-and-open")

                trigger.click()
                doc = Document()
                doc.add_paragraph("before table")
                table = doc.add_table(rows=1, cols=1)
                table.cell(0, 0).text = "inside table"
                doc.add_paragraph("after table")
                blob = io.BytesIO()
                doc.save(blob)
                choose(f"ordered-{width}.docx", blob.getvalue())
                expect(page.locator("#docs-import-save")).to_be_enabled()
                text = page.locator("#docs-import-preview").text_content()
                assert (
                    text.index("before table")
                    < text.index("inside table")
                    < text.index("after table")
                )
                page.locator("#docs-import-save").click()
                expect(page.locator("#docs-import-recovery .note-saved-open")).to_be_visible()
                page.locator("#docs-import-recovery .note-saved-open").click()
                expect(page.locator("#wiki-preview")).to_contain_text("after table")
                record("docx-ordered-import")

                if width == 390:
                    page.locator("#wiki-tree-toggle").click()
                page.locator("#wiki-search").fill(f"import searchable {width}")
                hit = page.locator("#wiki-tree .docs-search-hit").filter(
                    has_text=f"import-{width} 2"
                )
                expect(hit).to_be_visible()
                hit.click()
                expect(page.locator("#wiki-preview")).to_contain_text("edited after import")
                record("find-imported-content-through-ui")
                page.locator("#wiki-more-btn").click()
                page.locator("#wiki-delete-btn").click()
                page.locator('[data-dialog-action="delete"]').click()
                expect(page.locator("#wiki-document")).to_be_hidden()
                if width == 390:
                    page.locator("#wiki-tree-toggle").click()
                page.locator("#wiki-trash-btn").click()
                page.get_by_role("button", name=f"restore {path}", exact=True).click()
                expect(page.locator("#wiki-preview")).to_contain_text("edited after import")
                expect(page.locator("#wiki-document")).to_be_visible()
                expect(page.locator("#wiki-path")).to_have_text(path)
                record("trash-restore-imported-document")

                trigger.click()
                choose(f"focus-{width}.txt", b"focus stays where the owner moved it")
                expect(page.locator("#docs-import-save")).to_be_enabled()
                held = []

                def hold_save(route):
                    if route.request.method == "POST":
                        held.append((route, route.fetch()))
                    else:
                        route.continue_()

                page.route(base + "/api/vault-md/file", hold_save)
                page.locator("#docs-import-save").click()
                for _ in range(100):
                    if held:
                        break
                    page.wait_for_timeout(20)
                assert len(held) == 1
                trigger.focus()
                route, response = held[0]
                route.fulfill(response=response)
                expect(page.locator("#docs-import-panel")).to_be_hidden()
                expect(trigger).to_be_focused()
                page.unroute(base + "/api/vault-md/file", hold_save)
                record("late-save-preserves-moved-focus")
                trigger.click()
                choose(f"focused-preview-{width}.txt", b"preview keeps keyboard ownership")
                expect(page.locator("#docs-import-save")).to_be_enabled()
                held.clear()
                page.route(base + "/api/vault-md/file", hold_save)
                page.locator("#docs-import-save").click()
                for _ in range(100):
                    if held:
                        break
                    page.wait_for_timeout(20)
                assert len(held) == 1
                page.locator("#docs-import-preview").focus()
                expect(page.locator("#docs-import-preview")).to_be_focused()
                route, response = held[0]
                route.fulfill(response=response)
                expect(page.locator("#docs-import-panel")).to_be_hidden()
                expect(page.locator("#docs-import-recovery .note-saved-open")).to_be_focused()
                page.unroute(base + "/api/vault-md/file", hold_save)
                record("hidden-preview-transfers-keyboard-focus")

                assert page.evaluate("document.documentElement.scrollWidth <= innerWidth + 1")
                assert not errors, errors
                unexpected = [
                    line for line in console if not any(code in line for code in ("400", "503"))
                ]
                assert not unexpected, unexpected
                page.screenshot(path=str(out / f"{width}-import.png"), full_page=True)
                (out / f"{width}-console.json").write_text(
                    json.dumps({"page_errors": errors, "console": console}, indent=2)
                )
                record("keyboard-reduced-motion-no-overflow-console")
            except Exception as error:
                rows.append(
                    {
                        "scenario_id": "docs.import.current",
                        "profile": str(width),
                        "status": "failed",
                        "error": str(error),
                    }
                )
                page.screenshot(path=str(out / f"{width}-failed.png"), full_page=True)
                raise
            finally:
                (out / "scenarios.json").write_text(json.dumps(rows, indent=2))
                context.tracing.stop(path=str(out / f"{width}-import.zip"))
                context.close()
        api.dispose()
        browser.close()


if __name__ == "__main__":
    run()
